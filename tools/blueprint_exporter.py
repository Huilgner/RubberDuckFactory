#!/usr/bin/env python3
"""
blueprint_exporter.py — Exportador de Blueprints para DOCX e PDF

Converte os blueprints Markdown do Sovereign em documentos profissionais.

Uso:
    uv run python blueprint_exporter.py --projeto SIGO_FENIX --formato ambos
    uv run python blueprint_exporter.py --projeto CastleVote --formato docx
    uv run python blueprint_exporter.py --arquivo docs/blueprints/X/blueprint_R00.md --formato pdf
"""

import argparse
import re
import sys
from pathlib import Path

import os
from dotenv import load_dotenv

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT_DIR       = Path(__file__).parent

# Carrega configurações do .env se existir
load_dotenv(ROOT_DIR / ".env")
env_dir = os.getenv("BLUEPRINTS_DIR", "docs/blueprints")
BLUEPRINTS_DIR = ROOT_DIR / env_dir if not Path(env_dir).is_absolute() else Path(env_dir)

# ── Paleta ────────────────────────────────────────────────────────────────────

C_TITLE  = (0x1e, 0x3a, 0x5f)  # azul escuro
C_H2     = (0x18, 0x4e, 0x7a)  # azul
C_H3     = (0x33, 0x5a, 0x6a)  # teal-gray
C_META   = (0x64, 0x74, 0x8b)  # slate
C_SEP    = (0x94, 0xa3, 0xb8)  # cinza claro
C_TH_BG  = "#e2e8f0"
C_ALT_BG = "#f8fafc"


# ── Parsing Markdown ──────────────────────────────────────────────────────────

def split_bold(text: str):
    """Divide texto em lista de (is_bold, segment)."""
    result, last = [], 0
    for m in re.finditer(r"\*\*(.+?)\*\*", text):
        if m.start() > last:
            result.append((False, text[last:m.start()]))
        result.append((True, m.group(1)))
        last = m.end()
    if last < len(text):
        result.append((False, text[last:]))
    return result or [(False, text)]


def md_to_rl(text: str) -> str:
    """Converte **bold** para tags ReportLab <b>.</b>"""
    return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)


def parse_table(lines: list):
    headers, rows = [], []
    for i, line in enumerate(lines):
        cells = [c.strip() for c in line.strip("|").split("|")]
        if i == 0:
            headers = cells
        elif re.fullmatch(r"[\s\-|:]+", line.strip()):
            continue
        else:
            rows.append(cells)
    return headers, rows


def parse_lines(content: str) -> list:
    lines   = content.split("\n")
    elems   = []
    i       = 0
    sep_cnt = 0

    while i < len(lines):
        raw  = lines[i]
        line = raw.strip()

        if line.startswith("─") or line.startswith("---") and len(line) > 10:
            sep_cnt += 1
            elems.append({"type": "separator", "n": sep_cnt})

        elif line.startswith("BLUEPRINT") and "|" in line:
            parts = [p.strip() for p in line.split("|")]
            elems.append({"type": "bp_title", "parts": parts})

        elif line.startswith("<!-- sovereign-signature"):
            sig = []
            while i < len(lines):
                sig.append(lines[i].strip())
                if "-->" in lines[i]:
                    break
                i += 1
            elems.append({"type": "signature", "lines": sig})

        elif line.startswith("## "):
            elems.append({"type": "h2", "content": line[3:]})

        elif line.startswith("### "):
            elems.append({"type": "h3", "content": line[4:]})

        elif line.startswith("|") and "|" in line[1:]:
            tbl = [line]
            while i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
                i += 1
                tbl.append(lines[i].strip())
            h, r = parse_table(tbl)
            elems.append({"type": "table", "headers": h, "rows": r})

        elif line.startswith("- "):
            elems.append({"type": "bullet", "content": line[2:]})

        elif re.match(r"^[A-Z][A-Z ]{3,}\s{3,}", line):
            elems.append({"type": "meta_line", "content": line})

        elif not line:
            elems.append({"type": "empty"})

        else:
            elems.append({"type": "paragraph", "content": line})

        i += 1

    return elems


# ── DOCX ──────────────────────────────────────────────────────────────────────

def _cell_bg(cell, hex6: str):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex6)
    tcPr.append(shd)


def export_docx(elements: list, out: Path, projeto: str, rev: str):
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    for sec in doc.sections:
        sec.top_margin    = Cm(2)
        sec.bottom_margin = Cm(2)
        sec.left_margin   = Cm(2.5)
        sec.right_margin  = Cm(2.5)

    def rgb(*t):
        return RGBColor(*t)

    def para(style=None):
        return doc.add_paragraph(style=style)

    def add_runs(p, text: str, size: int, color, bold=False):
        for is_bold, seg in split_bold(text):
            r = p.add_run(seg)
            r.font.size  = Pt(size)
            r.font.color.rgb = rgb(*color)
            r.bold = bold or is_bold

    for elem in elements:
        t = elem["type"]

        if t == "separator":
            p = para()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after  = Pt(3)
            r = p.add_run("─" * 72)
            r.font.size      = Pt(8)
            r.font.color.rgb = rgb(*C_SEP)

        elif t == "bp_title":
            parts = elem["parts"]
            p = para()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after  = Pt(3)
            for j, part in enumerate(parts):
                if j:
                    sep = p.add_run("  |  ")
                    sep.font.size      = Pt(11)
                    sep.font.color.rgb = rgb(*C_SEP)
                r = p.add_run(part)
                r.font.size      = Pt(15 if j == 0 else 11)
                r.bold           = j in (0, 2)
                r.font.color.rgb = rgb(*C_TITLE)

        elif t == "meta_line":
            m = re.match(r"^([A-Z][A-Z ]+?)\s{2,}(.+)$", elem["content"])
            p = para()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after  = Pt(1)
            if m:
                rk = p.add_run(f"{m.group(1).strip():<16}")
                rk.font.size = Pt(10); rk.bold = True; rk.font.color.rgb = rgb(*C_META)
                rv = p.add_run(m.group(2).strip())
                rv.font.size = Pt(10); rv.font.color.rgb = rgb(*C_TITLE)
            else:
                add_runs(p, elem["content"], 10, C_META)

        elif t == "h2":
            p = para()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after  = Pt(4)
            r = p.add_run(elem["content"])
            r.font.size = Pt(13); r.bold = True; r.font.color.rgb = rgb(*C_H2)

        elif t == "h3":
            p = para()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after  = Pt(2)
            r = p.add_run(elem["content"])
            r.font.size = Pt(11); r.bold = True; r.font.color.rgb = rgb(*C_H3)

        elif t == "bullet":
            p = para("List Bullet")
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after  = Pt(1)
            add_runs(p, elem["content"], 10, (0, 0, 0))

        elif t == "table":
            headers = elem["headers"]
            rows    = elem["rows"]
            if not headers:
                continue
            all_rows = [headers] + rows
            n = len(headers)
            tbl = doc.add_table(rows=len(all_rows), cols=n)
            tbl.style = "Table Grid"
            for ri, row_data in enumerate(all_rows):
                for ci, cell_text in enumerate(row_data[:n]):
                    cell = tbl.cell(ri, ci)
                    cell.text = ""
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after  = Pt(2)
                    add_runs(p, cell_text, 9, (0, 0, 0), bold=(ri == 0))
                    if ri == 0:
                        _cell_bg(cell, "E2E8F0")
            col_w = Cm(15.5) // n
            for col in tbl.columns:
                for cell in col.cells:
                    cell.width = col_w
            doc.add_paragraph().paragraph_format.space_after = Pt(4)

        elif t == "signature":
            p = para()
            p.paragraph_format.space_before = Pt(10)
            r = p.add_run("─" * 40)
            r.font.size = Pt(8); r.font.color.rgb = rgb(*C_SEP)
            for line in elem["lines"]:
                if ":" in line and not line.startswith("<!--") and "-->" not in line:
                    p2 = para()
                    p2.paragraph_format.space_before = Pt(0)
                    p2.paragraph_format.space_after  = Pt(0)
                    rv = p2.add_run(line.strip())
                    rv.font.size = Pt(8); rv.font.color.rgb = rgb(*C_META)

        elif t == "paragraph":
            if not elem["content"]:
                continue
            p = para()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after  = Pt(2)
            add_runs(p, elem["content"], 10, (0, 0, 0))

    doc.save(str(out))
    print(f"  DOCX -> {out}")


# ── PDF ───────────────────────────────────────────────────────────────────────

def export_pdf(elements: list, out: Path, projeto: str, rev: str):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer,
        Table, TableStyle, HRFlowable,
    )
    from reportlab.lib.enums import TA_CENTER

    c_title  = colors.HexColor("#1e3a5f")
    c_h2     = colors.HexColor("#184e7a")
    c_h3     = colors.HexColor("#335a6a")
    c_meta   = colors.HexColor("#64748b")
    c_sep    = colors.HexColor("#cbd5e1")
    c_th_bg  = colors.HexColor(C_TH_BG)
    c_alt    = colors.HexColor(C_ALT_BG)

    def S(name, **kw):
        return ParagraphStyle(name, **kw)

    s_title  = S("T", fontSize=16, textColor=c_title, fontName="Helvetica-Bold",
                 alignment=TA_CENTER, spaceAfter=4, spaceBefore=6)
    s_meta   = S("M", fontSize=9,  textColor=c_meta,  fontName="Helvetica",
                 spaceAfter=2, spaceBefore=1)
    s_h2     = S("H2", fontSize=13, textColor=c_h2,   fontName="Helvetica-Bold",
                 spaceBefore=14, spaceAfter=4)
    s_h3     = S("H3", fontSize=11, textColor=c_h3,   fontName="Helvetica-Bold",
                 spaceBefore=8,  spaceAfter=2)
    s_body   = S("B", fontSize=10, textColor=colors.black, fontName="Helvetica",
                 spaceBefore=2, spaceAfter=2, leading=14)
    s_bullet = S("BL", fontSize=10, textColor=colors.black, fontName="Helvetica",
                 spaceBefore=1, spaceAfter=1, leftIndent=14, leading=13)
    s_sig    = S("SG", fontSize=8, textColor=c_meta,  fontName="Helvetica",
                 spaceBefore=1, spaceAfter=1)
    s_tc     = S("TC", fontSize=9, fontName="Helvetica", leading=12)
    s_th     = S("TH", fontSize=9, fontName="Helvetica-Bold", leading=12)

    story = []

    for elem in elements:
        t = elem["type"]

        if t == "separator":
            story.append(HRFlowable(width="100%", thickness=0.5, color=c_sep,
                                    spaceAfter=4, spaceBefore=4))

        elif t == "bp_title":
            parts = elem["parts"]
            safe  = [p.replace("&", "&amp;") for p in parts]
            line  = "  <font color='#94a3b8'>|</font>  ".join(
                f"<b>{p}</b>" if j in (0, 2) else p
                for j, p in enumerate(safe)
            )
            story.append(Paragraph(line, s_title))

        elif t == "meta_line":
            m = re.match(r"^([A-Z][A-Z ]+?)\s{2,}(.+)$", elem["content"])
            if m:
                story.append(Paragraph(
                    f"<b>{m.group(1).strip()}</b>&nbsp;&nbsp;&nbsp;&nbsp;{m.group(2).strip()}", s_meta))
            else:
                story.append(Paragraph(md_to_rl(elem["content"]), s_meta))

        elif t == "h2":
            story.append(Paragraph(md_to_rl(elem["content"]), s_h2))

        elif t == "h3":
            story.append(Paragraph(md_to_rl(elem["content"]), s_h3))

        elif t == "bullet":
            story.append(Paragraph(f"• {md_to_rl(elem['content'])}", s_bullet))

        elif t == "table":
            headers = elem["headers"]
            rows    = elem["rows"]
            if not headers:
                continue
            n       = len(headers)
            avail   = 16.5 * cm
            col_w   = [avail / n] * n

            data = [[Paragraph(md_to_rl(c), s_th) for c in headers]]
            for row in rows:
                padded = (row + [""] * n)[:n]
                data.append([Paragraph(md_to_rl(c), s_tc) for c in padded])

            tbl = Table(data, colWidths=col_w, repeatRows=1)
            tbl.setStyle(TableStyle([
                ("BACKGROUND",    (0, 0), (-1, 0),  c_th_bg),
                ("GRID",          (0, 0), (-1, -1), 0.5, c_sep),
                ("TOPPADDING",    (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("LEFTPADDING",   (0, 0), (-1, -1), 6),
                ("RIGHTPADDING",  (0, 0), (-1, -1), 6),
                ("VALIGN",        (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS",(0, 1), (-1, -1), [colors.white, c_alt]),
            ]))
            story.append(tbl)
            story.append(Spacer(1, 0.3 * cm))

        elif t == "empty":
            story.append(Spacer(1, 0.15 * cm))

        elif t == "signature":
            story.append(HRFlowable(width="100%", thickness=0.5, color=c_sep,
                                    spaceBefore=10, spaceAfter=4))
            for line in elem["lines"]:
                if ":" in line and not line.startswith("<!--") and "-->" not in line:
                    story.append(Paragraph(line.strip(), s_sig))

        elif t == "paragraph":
            if not elem["content"]:
                continue
            story.append(Paragraph(md_to_rl(elem["content"]), s_body))

    doc = SimpleDocTemplate(
        str(out), pagesize=A4,
        topMargin=2*cm, bottomMargin=2*cm,
        leftMargin=2.5*cm, rightMargin=2.5*cm,
        title=f"Blueprint {projeto} {rev}",
        author="Sovereign / RubberDuckFactory Squad",
    )
    doc.build(story)
    print(f"  PDF  -> {out}")


# ── Entry point ───────────────────────────────────────────────────────────────

def get_latest_blueprint(projeto: str):
    proj_dir   = BLUEPRINTS_DIR / projeto
    blueprints = sorted(proj_dir.glob("blueprint_R*.md"), reverse=True)
    if not blueprints:
        raise SystemExit(f"Nenhum blueprint encontrado para '{projeto}'")
    latest  = blueprints[0]
    rev     = re.search(r"(R\d+)", latest.name)
    rev_str = rev.group(1) if rev else "R00"
    return latest, latest.read_text(encoding="utf-8"), rev_str


def main():
    parser = argparse.ArgumentParser(
        description="Exportador de Blueprints — DOCX / PDF",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "Exemplos:\n"
            "  uv run python blueprint_exporter.py -p SIGO_FENIX\n"
            "  uv run python blueprint_exporter.py -p CastleVote -f pdf\n"
            "  uv run python blueprint_exporter.py -a docs/blueprints/X/blueprint_R01.md -f docx\n"
        ),
    )
    src = parser.add_mutually_exclusive_group(required=True)
    src.add_argument("--projeto", "-p", metavar="NOME")
    src.add_argument("--arquivo", "-a", metavar="PATH")

    parser.add_argument("--formato", "-f",
                        choices=["docx", "pdf", "ambos"], default="ambos")
    args = parser.parse_args()

    if args.projeto:
        md_path, content, rev = get_latest_blueprint(args.projeto)
        projeto = args.projeto
    else:
        md_path = Path(args.arquivo)
        if not md_path.exists():
            raise SystemExit(f"Arquivo nao encontrado: {md_path}")
        content = md_path.read_text(encoding="utf-8")
        rev_m   = re.search(r"(R\d+)", md_path.name)
        rev     = rev_m.group(1) if rev_m else "R00"
        projeto = md_path.parent.name

    elems   = parse_lines(content)
    out_dir = md_path.parent
    base    = md_path.stem

    print(f"\nExportando Blueprint {projeto} {rev} ...")
    if args.formato in ("docx", "ambos"):
        export_docx(elems, out_dir / f"{base}.docx", projeto, rev)
    if args.formato in ("pdf", "ambos"):
        export_pdf(elems,  out_dir / f"{base}.pdf",  projeto, rev)
    print(f"\nConcluido. Arquivos em: {out_dir}")


if __name__ == "__main__":
    main()
